from __future__ import annotations

import argparse
import os
import re
import sys
from pathlib import Path
from urllib.parse import parse_qs, urlparse

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from playwright.sync_api import Error as PlaywrightError
from playwright.sync_api import TimeoutError as PlaywrightTimeoutError
from playwright.sync_api import sync_playwright


def validate_url(url: str) -> None:
    parsed = urlparse(url)
    if parsed.scheme not in {"http", "https"} or parsed.hostname not in {
        "drafter.lol",
        "www.drafter.lol",
    } or not parsed.path.startswith("/draft/"):
        raise ValueError(f"Lien Drafter invalide : {url}")


def read_team_name(page, side: str) -> str:
    label = f"{side} side"
    element = page.get_by_text(label, exact=True).first

    for _ in range(6):
        text = re.sub(r"\s+", " ", element.inner_text()).strip()
        before_label, separator, _ = text.partition(label)
        if separator:
            name = re.sub(r"\s+(?:Ready|Not ready)\s*$", "", before_label, flags=re.I).strip()
            if name and name.casefold() != "ready":
                return name
        element = element.locator("xpath=..")

    return side.upper()


def read_champions(page, selector: str) -> list[str]:
    return [text.strip() for text in page.locator(selector).all_inner_texts()]


def read_draft(page, url: str) -> dict[str, object]:
    validate_url(url)
    try:
        response = page.goto(url, wait_until="domcontentloaded", timeout=60_000)
        page.locator(".group.blue-pick").first.wait_for(state="visible", timeout=45_000)
    except PlaywrightTimeoutError as error:
        raise RuntimeError(f"Le draft n'a pas chargé : {url}") from error
    except PlaywrightError as error:
        raise RuntimeError(f"Impossible d'ouvrir {url} dans Edge : {error}") from error

    if response is not None and response.status >= 400:
        raise RuntimeError(f"Drafter.lol a répondu HTTP {response.status} pour {url}")

    draft = {
        "blue_team": read_team_name(page, "blue"),
        "red_team": read_team_name(page, "red"),
        "blue_picks": read_champions(page, ".group.blue-pick"),
        "red_picks": read_champions(page, ".group.red-pick"),
        "blue_bans": read_champions(page, ".group.blue-ban"),
        "red_bans": read_champions(page, ".group.red-ban"),
        "game": parse_qs(urlparse(url).query).get("game", [""])[0],
    }

    for key in ("blue_picks", "red_picks", "blue_bans", "red_bans"):
        if len(draft[key]) != 5:
            raise RuntimeError(f"Cinq emplacements attendus pour {key} dans {url}")

    return draft


def draft_rows(draft: dict[str, object]) -> list[tuple[str, str, str, str]]:
    blue_picks = draft["blue_picks"]
    red_picks = draft["red_picks"]
    blue_bans = draft["blue_bans"]
    red_bans = draft["red_bans"]
    rows = []

    for index in range(3):
        rows.append((blue_bans[index], "BANS", "BANS", red_bans[index]))
    for index in range(3):
        number = index + 1
        rows.append((blue_picks[index], f"B{number}", f"R{number}", red_picks[index]))
    for index in range(3, 5):
        rows.append((blue_bans[index], "BANS", "BANS", red_bans[index]))
    for index in range(3, 5):
        number = index + 1
        rows.append((blue_picks[index], f"B{number}", f"R{number}", red_picks[index]))

    return rows


def add_draft_sheet(workbook: Workbook, draft: dict[str, object], number: int) -> None:
    sheet = workbook.create_sheet(title=f"Draft {number}")
    headers = ["BLUE", draft["blue_team"], draft["red_team"], "RED"]
    for column, value in enumerate(headers, start=1):
        sheet.cell(row=1, column=column, value=value)

    rows = draft_rows(draft)
    for row_index, values in enumerate(rows, start=2):
        for column, value in enumerate(values, start=1):
            sheet.cell(row=row_index, column=column, value=value)

    header_fills = ["1F4E78", "FFD966", "A9D18E", "C00000"]
    thin_black = Side(style="thin", color="000000")
    for column in range(1, 5):
        cell = sheet.cell(row=1, column=column)
        cell.fill = PatternFill("solid", fgColor=header_fills[column - 1])
        cell.font = Font(
            name="Calibri",
            size=11,
            bold=True,
            color="FFFFFF" if column in (1, 4) else "000000",
        )
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_index in range(2, len(rows) + 2):
        for column in range(1, 5):
            cell = sheet.cell(row=row_index, column=column)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(left=thin_black, right=thin_black, top=thin_black, bottom=thin_black)
            if column in (2, 3):
                cell.fill = PatternFill("solid", fgColor="000000")
                cell.font = Font(name="Calibri", size=11, color="FFFFFF", bold=True)

    widths = (22, max(18, min(30, len(str(headers[1])) + 3)), max(18, min(30, len(str(headers[2])) + 3)), 18)
    for column, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width

    sheet.row_dimensions[1].height = 30
    for row_index in range(2, len(rows) + 2):
        sheet.row_dimensions[row_index].height = 22
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet.print_area = f"A1:D{len(rows) + 1}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True


def create_workbook(urls: list[str], output: Path) -> None:
    with sync_playwright() as playwright:
        browser_channel = os.getenv("DRAFTER_BROWSER_CHANNEL", "msedge").strip()
        browser_name = "Chromium" if browser_channel.casefold() == "chromium" else "Microsoft Edge"
        launch_options = {
            "headless": False,
            "args": ["--disable-blink-features=AutomationControlled"],
        }
        try:
            if browser_channel.casefold() == "chromium":
                browser = playwright.chromium.launch(**launch_options)
            else:
                browser = playwright.chromium.launch(channel=browser_channel, **launch_options)
        except PlaywrightError as error:
            raise RuntimeError(f"{browser_name} est requis pour lire les drafts.") from error

        workbook = Workbook()
        workbook.remove(workbook.active)
        try:
            for number, url in enumerate(urls, start=1):
                page = browser.new_page(viewport={"width": 1440, "height": 1100})
                page.add_init_script(
                    "Object.defineProperty(navigator, 'webdriver', {get: () => undefined})"
                )
                try:
                    draft = read_draft(page, url)
                    add_draft_sheet(workbook, draft, number)
                finally:
                    page.close()
        finally:
            browser.close()

    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Extrait les picks et bans de liens Drafter.lol vers un classeur Excel."
    )
    parser.add_argument("urls", nargs="+", help="Un ou plusieurs liens /draft/ de Drafter.lol")
    parser.add_argument("-o", "--output", default="drafter_drafts.xlsx", help="Fichier Excel de sortie")
    args = parser.parse_args()

    try:
        create_workbook(args.urls, Path(args.output))
    except (ValueError, RuntimeError) as error:
        print(f"Erreur : {error}", file=sys.stderr)
        return 1

    print(f"Classeur créé : {args.output} ({len(args.urls)} draft(s))")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())